from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("resume/Miller-Emion-Resume-UPDATED.docx")


def set_spacing(paragraph, before=0, after=4, line=1.08):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_run(run, size=10.5, bold=False, color="111827", font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_rule(paragraph, color="D9E2DD"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_heading(doc, text):
    p = doc.add_paragraph()
    set_spacing(p, before=7, after=4, line=1)
    set_run(p.add_run(text.upper()), size=9, bold=True, color="075C45")
    add_rule(p)


def add_body(doc, text, after=4):
    p = doc.add_paragraph()
    set_spacing(p, after=after, line=1.10)
    set_run(p.add_run(text), size=9.6, color="263241")


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22 + level * 0.18)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    set_spacing(p, after=2, line=1.08)
    set_run(p.add_run(text), size=9.3, color="263241")


def add_role(doc, title, meta, bullets):
    p = doc.add_paragraph()
    set_spacing(p, before=2, after=0, line=1)
    set_run(p.add_run(title), size=10.3, bold=True, color="111827")
    p2 = doc.add_paragraph()
    set_spacing(p2, after=2, line=1)
    set_run(p2.add_run(meta), size=9.1, bold=True, color="526173")
    for bullet in bullets:
        add_bullet(doc, bullet)


def add_subrole(doc, title, dates, bullets):
    p = doc.add_paragraph()
    set_spacing(p, before=3, after=0, line=1)
    set_run(p.add_run(title), size=9.8, bold=True, color="111827")
    set_run(p.add_run(f" | {dates}"), size=9.1, bold=True, color="526173")
    for bullet in bullets:
        add_bullet(doc, bullet)


def add_project(doc, title, subtitle, bullets):
    add_role(doc, title, subtitle, bullets)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(9.6)

    title = doc.add_paragraph()
    set_spacing(title, after=0, line=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(title.add_run("Miller Emion"), size=20, bold=True, color="111827")

    subtitle = doc.add_paragraph()
    set_spacing(subtitle, after=1, line=1)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(subtitle.add_run("Junior Software Engineer / IT Support Specialist"), size=10.5, bold=True, color="075C45")

    contact = doc.add_paragraph()
    set_spacing(contact, after=5, line=1)
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(
        contact.add_run(
            "Cibolo, TX 78108 | DevDadx64@outlook.com | (719) 666-0177 | github.com/XtraSaltyDev | memion.io"
        ),
        size=8.7,
        color="526173",
    )

    add_heading(doc, "Summary")
    add_body(
        doc,
        "Junior software engineer and IT support specialist with experience progressing from frontline technical support into software engineering in a confidentiality-sensitive environment. Builds and maintains practical software while bringing a support-minded approach to troubleshooting, documentation, user communication, and operational reliability. Army veteran with CompTIA A+, CompTIA Network+, an Associate of Applied Science in Information Technology and Networking, and a completed Bachelor of Science in Software Design and Programming.",
        after=3,
    )

    add_heading(doc, "Technical Skills")
    for line in [
        "Languages and web: TypeScript, JavaScript, Python, PHP, Java, Swift, HTML, CSS",
        "Developer tools: Git, GitHub, Node.js, npm, Astro, Vercel, SwiftPM, Vitest, ESLint, Prettier, VS Code, PhpStorm",
        "IT support: Active Directory, Remote Desktop, Splashtop, Duo Security, Microsoft 365, Microsoft Teams, WinSCP, PuTTY, Wireshark",
        "Systems: Windows, macOS, Linux, local development environments, troubleshooting, documentation, SOP creation",
    ]:
        add_bullet(doc, line)

    add_heading(doc, "Projects")
    add_project(
        doc,
        "agent-notes",
        "Public TypeScript CLI / npm package for deterministic repository context.",
        [
            "Built and published a CLI that scans Node.js and TypeScript repositories and generates reviewable AGENTS.md plus .agent-notes/ Markdown files.",
            "Implemented scan, init, update, and doctor workflows with JSON output, dry-run support, path targeting, and marker-managed updates.",
            "Used TypeScript, Commander, Vitest, ESLint, Prettier, tsup, npm packaging checks, tarball smoke tests, and fresh-cache npx verification.",
        ],
    )
    add_project(
        doc,
        "memion.io",
        "Personal portfolio site built with Astro and deployed through Vercel.",
        [
            "Built a responsive static homepage with plain CSS, Astro, metadata, navigation anchors, and project/contact sections.",
            "Refactored the page information architecture so identity, work, notes, and contact content each have one clear location.",
            "Configured production deployment path with Vercel and Cloudflare-managed domain routing for memion.io.",
        ],
    )
    add_project(
        doc,
        "Open WebUI Native Experiment",
        "SwiftPM macOS project exploring native Open WebUI-style workflows.",
        [
            "Worked in a Swift/macOS codebase with models, services, and views for chat, settings, knowledge, tools, automation, audio, code execution, and local storage.",
            "Practiced building JSON-backed storage/export services, provider models, and SwiftUI-style application surfaces.",
            "Focused on preserving useful Open WebUI concepts while learning native app structure and local-first workflows.",
        ],
    )

    add_heading(doc, "Experience")
    p = doc.add_paragraph()
    set_spacing(p, before=2, after=0, line=1)
    set_run(p.add_run("Gavin de Becker & Associates"), size=10.3, bold=True, color="111827")
    p2 = doc.add_paragraph()
    set_spacing(p2, after=1, line=1)
    set_run(p2.add_run("San Antonio, TX"), size=9.1, bold=True, color="526173")
    add_subrole(
        doc,
        "Junior Software Engineer",
        "January 2026 - Present",
        [
            "Build, test, and maintain internal software in a confidentiality-sensitive environment.",
            "Translate operational needs and support feedback into maintainable software improvements.",
            "Troubleshoot defects, document behavior, and collaborate across technical and operational teams while protecting sensitive information.",
        ],
    )
    add_subrole(
        doc,
        "Support Specialist 2",
        "May 2025 - January 2026",
        [
            "Handled escalated technical support issues across user access, endpoint, application, and workflow problems.",
            "Supported troubleshooting and documentation for internal systems while coordinating with engineering and systems teams.",
            "Improved support handoffs by communicating clearly, documenting repeatable fixes, and maintaining confidentiality.",
        ],
    )
    add_subrole(
        doc,
        "Support Specialist 1",
        "February 2024 - May 2025",
        [
            "Provided frontline technical support for users, devices, accounts, applications, and routine access issues.",
            "Documented issues, tracked follow-up, and escalated problems with clear context for senior support and engineering teams.",
            "Built a strong operational foundation in troubleshooting, user communication, remote support, and support-process discipline.",
        ],
    )
    add_role(
        doc,
        "Walmart, Inc.",
        "Asset Protection Investigator | Cibolo, TX | November 2022 - July 2023",
        [
            "Provided customer-facing support while handling inquiries, resolving conflicts, and documenting incidents.",
            "Collaborated with team members and other departments to maintain a safe, secure environment.",
            "Used computer systems and reporting tools for monitoring, documentation, and follow-up.",
            "Built multitasking and time-management habits in a fast-paced operational environment.",
        ],
    )
    add_role(
        doc,
        "United States Army",
        "Specialist, Combat Engineer | March 2019 - February 2022",
        [
            "Served in high-accountability team environments requiring communication, discipline, safety awareness, and reliable execution.",
            "Developed practical leadership, problem-solving, and operational planning habits under pressure.",
        ],
    )

    add_heading(doc, "Education")
    add_role(
        doc,
        "DeVry University",
        "Naperville, IL",
        [
            "Bachelor of Science in Software Design and Programming, October 2025.",
            "Associate of Applied Science in Information Technology and Networking, May 2024.",
            "Undergraduate Certificate in Programming Essentials, February 2024.",
            "Relevant coursework: operating systems, programming, digital devices, networking, information security, Linux administration, cybersecurity, virtualization, and cloud computing.",
        ],
    )

    add_heading(doc, "Certifications")
    add_bullet(doc, "CompTIA A+, April 2024")
    add_bullet(doc, "CompTIA Network+, August 2025")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
